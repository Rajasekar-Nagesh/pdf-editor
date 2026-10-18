"""Accessibility checks and auto-tagging, print preflight, and scanned tables to Excel."""
import re
from collections import Counter
from pathlib import Path

import pikepdf
import pymupdf as fitz
from openpyxl import Workbook

from .common import ToolContext, ToolError, ToolOutput, open_pdf

# ------------------------------------------------------------------ accessibility


def _item(id_, title, status, detail):
    return {"id": id_, "title": title, "status": status, "detail": detail}


def _struct_elements(node, depth=0, limit=20000):
    """Yield structure elements (dicts with /S) from the structure tree, iteratively bounded."""
    stack = [node]
    seen = 0
    while stack and seen < limit:
        cur = stack.pop()
        seen += 1
        if isinstance(cur, pikepdf.Array):
            stack.extend(list(cur))
            continue
        if not isinstance(cur, pikepdf.Dictionary):
            continue
        if "/S" in cur:
            yield cur
        kids = cur.get("/K")
        if kids is not None:
            stack.append(kids)


def accessibility_check(ctx: ToolContext) -> ToolOutput:
    items = []
    with pikepdf.open(ctx.inputs[0]) as pdf:
        root = pdf.Root
        marked = bool(root.get("/MarkInfo", {}).get("/Marked", False)) if "/MarkInfo" in root else False
        tree = root.get("/StructTreeRoot")
        items.append(_item("tagged", "Tagged PDF", "pass" if marked and tree is not None else "fail",
                           "Document is tagged." if tree is not None else "No structure tags. Screen readers cannot read it reliably. Use Auto-tag."))
        lang = str(root.get("/Lang", "")) if "/Lang" in root else ""
        items.append(_item("lang", "Document language", "pass" if lang else "fail", f"Language: {lang}" if lang else "No document language set."))
        title = ""
        with fitz.open(ctx.inputs[0]) as d:
            title = (d.metadata or {}).get("title", "")
        display_title = bool(root.get("/ViewerPreferences", {}).get("/DisplayDocTitle", False)) if "/ViewerPreferences" in root else False
        items.append(_item("title", "Document title", "pass" if title and display_title else ("warn" if title else "fail"),
                           f"Title: {title}" if title else "No title in document properties."))

        figures = []
        if tree is not None:
            figures = [el for el in _struct_elements(tree) if el.get("/S") == pikepdf.Name.Figure]
        missing_alt = [f for f in figures if "/Alt" not in f]
        image_count = 0
        for page in pdf.pages:
            res = page.obj.get("/Resources", {})
            xobj = res.get("/XObject", {}) if res is not None else {}
            for _k, x in (xobj.items() if xobj is not None else []):
                if x.get("/Subtype") == pikepdf.Name.Image:
                    image_count += 1
        if image_count == 0:
            items.append(_item("alt", "Alternate text for images", "pass", "No images found."))
        elif tree is None or not figures:
            items.append(_item("alt", "Alternate text for images", "fail", f"{image_count} image(s) are not tagged as figures with alt text."))
        else:
            items.append(_item("alt", "Alternate text for images", "fail" if missing_alt else "pass",
                               f"{len(missing_alt)} of {len(figures)} figure(s) have no alt text." if missing_alt else "All figures have alt text."))

        tabs_ok = all(page.obj.get("/Tabs") == pikepdf.Name.S for page in pdf.pages if "/Annots" in page.obj)
        items.append(_item("tabs", "Tab order follows structure", "pass" if tabs_ok else "warn",
                           "OK" if tabs_ok else "Pages with links/fields should use structure tab order."))
        parent_tree = tree is not None and "/ParentTree" in tree
        items.append(_item("order", "Logical reading order", "pass" if parent_tree else "fail",
                           "Reading order is defined by the structure tree." if parent_tree else "No reading order is defined (untagged content)."))
        page_count = len(pdf.pages)

    with fitz.open(ctx.inputs[0]) as d:
        text_chars = sum(len(p.get_text("text").strip()) for p in d)
        items.append(_item("text", "Real text (not scanned images)", "pass" if text_chars > 20 * max(1, d.page_count) else "fail",
                           "Document contains text." if text_chars else "No text layer. Run OCR first."))
        bad_fonts = set()
        for p in d:
            for f in p.get_fonts(full=True):
                if f[1] in ("Type3",):
                    bad_fonts.add(f[3])
        items.append(_item("fonts", "Fonts map to Unicode", "warn" if bad_fonts else "pass",
                           f"Type3 fonts may not map to text: {', '.join(sorted(bad_fonts))[:200]}" if bad_fonts else "OK"))
        has_toc = bool(d.get_toc())
    items.append(_item("bookmarks", "Bookmarks for long documents", "pass" if has_toc or page_count <= 20 else "warn",
                       "OK" if has_toc or page_count <= 20 else "Documents over 20 pages should have bookmarks."))
    failed = sum(1 for i in items if i["status"] == "fail")
    return ToolOutput(data={"items": items, "failed": failed, "passed": sum(1 for i in items if i["status"] == "pass")})


def _headings(doc: fitz.Document, max_items: int = 200) -> list[list]:
    """Guess headings from font sizes: noticeably larger than the body text, short lines."""
    sizes: Counter = Counter()
    lines = []
    for page in doc:
        for b in page.get_text("dict")["blocks"]:
            for line in b.get("lines", []):
                text = "".join(s["text"] for s in line["spans"]).strip()
                if not text:
                    continue
                size = max(s["size"] for s in line["spans"])
                sizes[round(size)] += len(text)
                lines.append((page.number, size, text, any(s["flags"] & 16 for s in line["spans"])))
    if not sizes:
        return []
    body = sizes.most_common(1)[0][0]
    big = sorted({round(s) for _, s, _, _ in lines if s >= body * 1.25}, reverse=True)[:3]
    toc = []
    for pno, size, text, _bold in lines:
        if round(size) in big and 2 < len(text) < 120:
            toc.append([big.index(round(size)) + 1, text, pno + 1])
            if len(toc) >= max_items:
                break
    # Levels must start at 1 and may only increase by one.
    fixed, prev = [], 0
    for lvl, text, page in toc:
        lvl = min(lvl, prev + 1)
        fixed.append([lvl, text, page])
        prev = lvl
    return fixed


def auto_tag(ctx: ToolContext, language: str = "en-US", title: str = "", alt_text: str = "", add_bookmarks: bool = True) -> ToolOutput:
    """Add a basic structure tree (one tagged section per page, in page order), language, title,
    tab order, and bookmarks generated from headings. Figures get the supplied alt text.
    """
    if not re.fullmatch(r"[A-Za-z]{2,3}(-[A-Za-z0-9]{2,8})*", language or ""):
        raise ToolError("Language must be a code like en-US, ta-IN or hi-IN.")
    mid = ctx.workdir / "titled.pdf"
    doc = open_pdf(ctx.inputs[0])
    meta = dict(doc.metadata or {})
    meta["title"] = title or meta.get("title") or ctx.stem
    doc.set_metadata(meta)
    if add_bookmarks and not doc.get_toc():
        toc = _headings(doc)
        if toc:
            doc.set_toc(toc)
    doc.save(mid, garbage=3, deflate=True)
    doc.close()

    name = f"{ctx.stem}-tagged.pdf"
    with pikepdf.open(mid) as pdf:
        root = pdf.Root
        if "/StructTreeRoot" in root:
            # Keep an existing (probably better) structure; only fix document-level properties.
            pass
        else:
            doc_elem = pdf.make_indirect(pikepdf.Dictionary(Type=pikepdf.Name.StructElem, S=pikepdf.Name.Document, K=pikepdf.Array()))
            nums = pikepdf.Array()
            struct_root = pdf.make_indirect(
                pikepdf.Dictionary(Type=pikepdf.Name.StructTreeRoot, K=doc_elem, ParentTree=pikepdf.Dictionary(Nums=nums), ParentTreeNextKey=len(pdf.pages))
            )
            doc_elem.P = struct_root
            for idx, page in enumerate(pdf.pages):
                has_images = any(
                    x.get("/Subtype") == pikepdf.Name.Image
                    for _k, x in (page.obj.get("/Resources", {}).get("/XObject", {}) or {}).items()
                ) if "/Resources" in page.obj else False
                text_elem = pdf.make_indirect(
                    pikepdf.Dictionary(Type=pikepdf.Name.StructElem, S=pikepdf.Name.Sect, P=doc_elem, Pg=page.obj, K=0)
                )
                doc_elem.K.append(text_elem)
                if has_images and alt_text:
                    fig = pikepdf.Dictionary(Type=pikepdf.Name.StructElem, S=pikepdf.Name.Figure, P=text_elem, Alt=pikepdf.String(alt_text))
                    text_elem.K = pikepdf.Array([0, pdf.make_indirect(fig)])
                page.obj.StructParents = idx
                nums.append(idx)
                nums.append(pikepdf.Array([text_elem]))
                page.contents_add(pikepdf.Stream(pdf, b"/Sect <</MCID 0>> BDC\n"), prepend=True)
                page.contents_add(pikepdf.Stream(pdf, b"\nEMC\n"), prepend=False)
                ctx.progress(int((idx + 1) * 90 / len(pdf.pages)), "Tagging")
            root.StructTreeRoot = struct_root
        root.MarkInfo = pikepdf.Dictionary(Marked=True)
        root.Lang = pikepdf.String(language)
        vp = root.get("/ViewerPreferences", pikepdf.Dictionary())
        vp.DisplayDocTitle = True
        root.ViewerPreferences = vp
        for page in pdf.pages:
            if "/Annots" in page.obj:
                page.obj.Tabs = pikepdf.Name.S
        pdf.save(ctx.out(name))
    return ToolOutput(files=[(ctx.out(name), name)])


# ------------------------------------------------------------------ print preflight


def preflight(ctx: ToolContext, min_dpi: int = 300, require_cmyk: bool = False) -> ToolOutput:
    doc = open_pdf(ctx.inputs[0])
    fonts: dict[str, dict] = {}
    images = []
    spot_colors = set()
    sizes = Counter()
    for page in doc:
        sizes[(round(page.rect.width), round(page.rect.height))] += 1
        for xref, ext, ftype, basefont, _name, _enc, *_ in page.get_fonts(full=True):
            fonts.setdefault(basefont, {"name": basefont, "type": ftype, "embedded": ext not in ("n/a", ""), "pages": []})
            if page.number + 1 not in fonts[basefont]["pages"]:
                fonts[basefont]["pages"].append(page.number + 1)
        for info in page.get_image_info(xrefs=True):
            bbox = fitz.Rect(info["bbox"])
            if bbox.width <= 0 or bbox.height <= 0:
                continue
            dpi_x = info["width"] / (bbox.width / 72)
            dpi_y = info["height"] / (bbox.height / 72)
            cs = info.get("cs-name") or ""
            images.append(
                {
                    "page": page.number + 1,
                    "xref": info.get("xref"),
                    "dpi": round(min(dpi_x, dpi_y)),
                    "colorspace": cs,
                    "low_resolution": min(dpi_x, dpi_y) < int(min_dpi),
                    "rgb_in_cmyk_job": bool(require_cmyk and "RGB" in cs.upper()),
                }
            )
        ctx.progress(int((page.number + 1) * 80 / doc.page_count), f"Checking page {page.number + 1}")
    for xref in range(1, doc.xref_length()):
        try:
            obj = doc.xref_object(xref, compressed=False)
        except Exception:
            continue
        for m in re.finditer(r"/Separation\s*/([^\s/\]]+)", obj):
            if m.group(1) not in ("All", "None"):
                spot_colors.add(m.group(1).replace("#20", " "))
    trim = any("/TrimBox" in doc.xref_object(p.xref) for p in doc)
    bleed = any("/BleedBox" in doc.xref_object(p.xref) for p in doc)
    not_embedded = [f for f in fonts.values() if not f["embedded"]]
    low_res = [i for i in images if i["low_resolution"]]
    rgb = [i for i in images if i["rgb_in_cmyk_job"]]
    checks = [
        _item("fonts", "All fonts embedded", "fail" if not_embedded else "pass",
              f"{len(not_embedded)} font(s) not embedded: {', '.join(f['name'] for f in not_embedded)[:200]}" if not_embedded else f"{len(fonts)} font(s), all embedded."),
        _item("resolution", f"Images at least {min_dpi} dpi", "warn" if low_res else "pass",
              f"{len(low_res)} image(s) below {min_dpi} dpi." if low_res else f"{len(images)} image(s) OK."),
        _item("color", "Colour space", "warn" if rgb else "pass",
              f"{len(rgb)} RGB image(s) in a CMYK print job." if rgb else "OK"),
        _item("spot", "Spot colours", "info" if spot_colors else "pass", ", ".join(sorted(spot_colors)) or "None"),
        _item("sizes", "Consistent page size", "pass" if len(sizes) == 1 else "warn",
              ", ".join(f"{w}x{h} pt ({n} pages)" for (w, h), n in sizes.items())),
        _item("boxes", "Trim / bleed boxes", "pass" if trim and bleed else "info",
              f"TrimBox: {'yes' if trim else 'no'}, BleedBox: {'yes' if bleed else 'no'}"),
    ]
    return ToolOutput(data={"checks": checks, "fonts": list(fonts.values()), "images": images[:500]})


# ------------------------------------------------------------------ scanned tables -> Excel


def _cluster_rows(words, tolerance: float) -> list[list]:
    rows: list[list] = []
    for w in sorted(words, key=lambda w: ((w[1] + w[3]) / 2, w[0])):
        mid = (w[1] + w[3]) / 2
        if rows and abs(rows[-1][0] - mid) <= tolerance:
            rows[-1][1].append(w)
        else:
            rows.append([mid, [w]])
    return [sorted(r[1], key=lambda w: w[0]) for r in rows]


def _words_to_table(words) -> list[list[str]]:
    """Group OCR words into rows, then into columns using gaps shared across rows."""
    if not words:
        return []
    heights = sorted(w[3] - w[1] for w in words)
    h = heights[len(heights) // 2] or 10
    rows = _cluster_rows(words, h * 0.5)
    # Merge words separated by a small gap into cells.
    cell_rows = []
    for row in rows:
        cells = []
        for w in row:
            if cells and w[0] - cells[-1][2] < h * 0.9:
                cells[-1] = [cells[-1][0], cells[-1][1], w[2], cells[-1][3] + " " + w[4]]
            else:
                cells.append([w[0], w[1], w[2], w[4]])
        cell_rows.append(cells)
    # Column anchors: cluster left edges of cells across all rows.
    lefts = sorted(c[0] for r in cell_rows for c in r)
    anchors: list[float] = []
    for x in lefts:
        if not anchors or x - anchors[-1] > h * 1.5:
            anchors.append(x)
    table = []
    for r in cell_rows:
        line = [""] * len(anchors)
        for c in r:
            idx = max(i for i, a in enumerate(anchors) if a <= c[0] + h * 0.75) if any(a <= c[0] + h * 0.75 for a in anchors) else 0
            line[idx] = (line[idx] + " " + c[3]).strip()
        table.append(line)
    # Drop empty columns.
    keep = [i for i in range(len(anchors)) if any(row[i] for row in table)]
    return [[row[i] for i in keep] for row in table]


def scanned_tables_to_excel(ctx: ToolContext, languages: list[str] | None = None) -> ToolOutput:
    """OCR a scanned invoice/bill, then rebuild its tables as spreadsheet rows."""
    from .convert import _tables_from_pdf, _write_sheet
    from .optimize import ocr_file

    src = ctx.inputs[0]
    with fitz.open(src) as d:
        if not d.is_pdf:
            pdf_path = ctx.workdir / "image.pdf"
            fitz.open("pdf", d.convert_to_pdf()).save(pdf_path)
            src = pdf_path
    ocred = ctx.workdir / "ocr.pdf"
    ctx.progress(5, "Recognising text")
    ocr_file(src, ocred, languages or ["eng"], mode="force", deskew=True)
    ctx.progress(60, "Finding tables")
    doc = fitz.open(ocred)
    wb = Workbook()
    wb.remove(wb.active)
    tables = _tables_from_pdf(Path(ocred), list(range(doc.page_count)), lambda *_: None)
    good = [(p, rows) for p, rows in tables if max(len(r) for r in rows) >= 3]
    if good:
        for k, (page, rows) in enumerate(good, 1):
            _write_sheet(wb.create_sheet(f"Table {k} (p{page + 1})"[:31]), rows)
    else:
        for page in doc:
            rows = _words_to_table(page.get_text("words"))
            if rows:
                _write_sheet(wb.create_sheet(f"Page {page.number + 1}"), rows)
    if not wb.sheetnames:
        raise ToolError("No text could be recognised in this scan. Try a clearer, higher-resolution scan.")
    name = f"{ctx.stem}.xlsx"
    wb.save(ctx.out(name))
    return ToolOutput(files=[(ctx.out(name), name)], data={"sheets": wb.sheetnames})
