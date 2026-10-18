import io
import zipfile

import pikepdf
import pymupdf as fitz
import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from .conftest import make_pdf, make_png, need


def _doc(fetch, headers, job, idx=0):
    return fitz.open(stream=fetch(headers, job["result"]["files"][idx]["id"]))


def _table_pdf() -> bytes:
    doc = fitz.open()
    p = doc.new_page()
    rows = [["Item", "Qty", "Price"], ["Pens", "10", "1,250.50"], ["Paper", "5", "300"]]
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            rect = fitz.Rect(72 + c * 120, 100 + r * 25, 192 + c * 120, 125 + r * 25)
            p.draw_rect(rect, color=(0, 0, 0))
            p.insert_text((rect.x0 + 5, rect.y1 - 8), val, fontsize=11)
    return doc.tobytes()


def test_pdf_to_word_excel_pptx(client, guest, upload, run, fetch):
    f = upload(guest)
    job = run(guest, "pdf_to_word", [f["id"]])
    docx = Document(io.BytesIO(fetch(guest, job["result"]["files"][0]["id"])))
    assert any("Hello PDFDesk" in p.text for p in docx.paragraphs)

    t = upload(guest, _table_pdf())
    job = run(guest, "pdf_to_excel", [t["id"]])
    assert job["result"]["data"]["tables"] >= 1
    wb = load_workbook(io.BytesIO(fetch(guest, job["result"]["files"][0]["id"])))
    values = [c.value for row in wb.worksheets[0].iter_rows() for c in row]
    assert "Pens" in values and 1250.5 in values  # numbers become real numbers

    job = run(guest, "pdf_to_pptx", [f["id"]], {"dpi": 72})
    assert len(Presentation(io.BytesIO(fetch(guest, job["result"]["files"][0]["id"]))).slides) == 3


def test_pdf_to_images_html_text(client, guest, upload, run, fetch):
    f = upload(guest)
    job = run(guest, "pdf_to_images", [f["id"]], {"fmt": "jpg", "dpi": 50})
    assert len(job["result"]["files"]) == 3
    with zipfile.ZipFile(io.BytesIO(fetch(guest, job["result"]["zip"]["id"]))) as z:
        assert len(z.namelist()) == 3
    job = run(guest, "pdf_to_images", [f["id"]], {"fmt": "png", "pages": "1"})
    assert fetch(guest, job["result"]["files"][0]["id"]).startswith(b"\x89PNG")
    assert b"Hello PDFDesk" in fetch(guest, run(guest, "pdf_to_html", [f["id"]])["result"]["files"][0]["id"])
    assert b"----- Page 2 -----" in fetch(guest, run(guest, "pdf_to_text", [f["id"]])["result"]["files"][0]["id"])


def test_images_to_pdf(client, guest, upload, run, fetch):
    a = upload(guest, make_png(), "a.png", "image/png")
    b = upload(guest, make_png((200, 0, 0)), "b.png", "image/png")
    assert _doc(fetch, guest, run(guest, "images_to_pdf", [a["id"], b["id"]], {"page_size": "a4"})).page_count == 2
    pdf = upload(guest)
    assert client.post("/api/tools/images_to_pdf", headers=guest, json={"file_ids": [pdf["id"]]}).status_code == 400


@need("soffice")
def test_office_to_pdf(client, guest, upload, run, fetch):
    buf = io.BytesIO()
    d = Document()
    d.add_paragraph("Quarterly report")
    d.save(buf)
    f = upload(guest, buf.getvalue(), "report.docx")
    job = run(guest, "office_to_pdf", [f["id"]])
    assert job["result"]["files"][0]["name"] == "report.pdf"
    assert "Quarterly report" in _doc(fetch, guest, job)[0].get_text()


def test_url_to_pdf_blocks_private_addresses(client, guest, run):
    for url in ("http://127.0.0.1:8000/", "http://localhost/", "file:///etc/passwd", "http://169.254.169.254/latest/meta-data"):
        job = run(guest, "url_to_pdf", [], {"url": url}, expect="failed")
        assert "private" in job["error"] or "http" in job["error"]


@need("gs")
def test_compress(client, guest, upload, run):
    f = upload(guest)
    job = run(guest, "compress", [f["id"]], {"quality": "low"})
    data = job["result"]["data"]
    assert data["after"] <= data["before"]


def test_compress_validation(client, guest, upload, run):
    f = upload(guest)
    assert run(guest, "compress", [f["id"]], {"quality": "ultra"}, expect="failed")["error"]


@need("gs")
def test_pdfa(client, guest, upload, run):
    f = upload(guest)
    job = run(guest, "to_pdfa", [f["id"]])
    assert "validation" in job["result"]["data"]


def test_validate_pdfa_without_verapdf(client, guest, upload, run):
    f = upload(guest)
    job = run(guest, "validate_pdfa", [f["id"]], {"flavour": "2b"})
    assert "available" in job["result"]["data"]["validation"]


@need("ocrmypdf")
def test_ocr(client, guest, upload, run, fetch):
    # Render text to an image-only PDF, then OCR it.
    src = fitz.open(stream=make_pdf(1))
    pix = src[0].get_pixmap(dpi=200)
    scan = fitz.open()
    page = scan.new_page(width=595, height=842)
    page.insert_image(page.rect, pixmap=pix)
    f = upload(guest, scan.tobytes())
    job = run(guest, "ocr", [f["id"]], {"languages": ["eng"]})
    assert "PDFDesk" in _doc(fetch, guest, job)[0].get_text()


@need("ocrmypdf")
def test_scanned_tables_to_excel(client, guest, upload, run):
    src = fitz.open(stream=_table_pdf())
    pix = src[0].get_pixmap(dpi=200)
    f = upload(guest, pix.tobytes("png"), "bill.png", "image/png")
    job = run(guest, "scanned_tables_to_excel", [f["id"]])
    assert job["result"]["data"]["sheets"]


def test_password_add_remove(client, guest, upload, run, fetch):
    f = upload(guest)
    job = run(guest, "add_password", [f["id"]], {"user_password": "s3cret", "allow_print": True, "allow_copy": False})
    protected = job["result"]["files"][0]["id"]
    with pytest.raises(pikepdf.PasswordError):
        pikepdf.open(io.BytesIO(fetch(guest, protected)))
    # The password was never stored in the job record.
    from app.db import SessionLocal
    from app.models import Job

    with SessionLocal() as db:
        assert db.get(Job, job["id"]).params["user_password"] == "***"

    assert "incorrect" in run(guest, "remove_password", [protected], {"password": "wrong"}, expect="failed")["error"]
    job = run(guest, "remove_password", [protected], {"password": "s3cret"})
    assert not pikepdf.open(io.BytesIO(fetch(guest, job["result"]["files"][0]["id"]))).is_encrypted
    assert run(guest, "add_password", [f["id"]], {}, expect="failed")["error"]


def test_redaction(client, guest, upload, run, fetch):
    f = upload(guest)
    job = run(guest, "find_sensitive", [f["id"]], {"presets": ["pan", "email", "phone"]})
    kinds = {m["kind"] for m in job["result"]["data"]["matches"]}
    assert {"PAN number", "Email address", "Phone number"} <= kinds

    job = run(guest, "redact", [f["id"]], {"presets": ["pan", "email"], "terms": ["Hello"], "areas": [{"page": 1, "rect": [400, 400, 480, 480]}]})
    d = _doc(fetch, guest, job)
    text = "".join(p.get_text() for p in d)
    assert "ABCDE1234F" not in text and "rajesh@example.com" not in text and "Hello" not in text
    assert all(abs(i["bbox"][0] - 400) > 1 for i in d[1].get_image_info())  # image under the area removed
    assert run(guest, "redact", [f["id"]], {}, expect="failed")["error"]
    assert "regular expression" in run(guest, "redact", [f["id"]], {"regex": "("}, expect="failed")["error"]


def test_sanitize(client, guest, upload, run, fetch):
    doc = fitz.open(stream=make_pdf(1))
    doc.set_metadata({"author": "Secret Author", "title": "Internal"})
    doc.embfile_add("notes.txt", b"hidden notes")
    p = doc[0]
    p.add_text_annot((100, 100), "internal comment")
    data = doc.tobytes()
    with pikepdf.open(io.BytesIO(data)) as pdf:
        pdf.Root.OpenAction = pikepdf.Dictionary(S=pikepdf.Name.JavaScript, JS=pikepdf.String("app.alert(1)"))
        buf = io.BytesIO()
        pdf.save(buf)
    f = upload(guest, buf.getvalue())
    job = run(guest, "sanitize", [f["id"]])
    assert job["result"]["data"]["javascript_entries_removed"] >= 0
    raw = fetch(guest, job["result"]["files"][0]["id"])
    d = fitz.open(stream=raw)
    assert d.embfile_count() == 0 and not d.metadata.get("author") and not list(d[0].annots())
    assert b"app.alert" not in raw
