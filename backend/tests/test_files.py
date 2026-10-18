import io
import uuid
from datetime import datetime, timedelta, timezone

import pymupdf as fitz
from docx import Document
from openpyxl import Workbook

from app.db import SessionLocal
from app.models import StoredFile
from app.scheduler import cleanup_expired_files

from .conftest import make_pdf, make_png


def test_upload_pdf_and_image(client, guest, upload):
    f = upload(guest)
    assert f["kind"] == "pdf" and f["expires_at"] is not None  # guest files expire
    img = upload(guest, make_png(), "photo.png", "image/png")
    assert img["kind"] == "image"


def test_upload_office_files(client, guest, upload):
    buf = io.BytesIO()
    d = Document()
    d.add_paragraph("Hello")
    d.save(buf)
    assert upload(guest, buf.getvalue(), "a.docx")["kind"] == "docx"
    buf = io.BytesIO()
    Workbook().save(buf)
    assert upload(guest, buf.getvalue(), "a.xlsx")["kind"] == "xlsx"


def test_upload_rejects_by_content(client, guest):
    # A text file renamed to .pdf is rejected because the content is checked, not the name.
    r = client.post("/api/files", headers=guest, files={"files": ("fake.pdf", b"just some text", "application/pdf")})
    assert r.status_code == 415
    r = client.post("/api/files", headers=guest, files={"files": ("x.exe", b"MZ\x90\x00binary", "application/octet-stream")})
    assert r.status_code == 415


def test_upload_size_limit(client, guest, monkeypatch):
    from app.config import get_settings

    monkeypatch.setattr(get_settings(), "max_upload_mb", 0)
    r = client.post("/api/files", headers=guest, files={"files": ("a.pdf", make_pdf(1), "application/pdf")})
    assert r.status_code == 413


def test_upload_requires_identity(client):
    assert client.post("/api/files", files={"files": ("a.pdf", make_pdf(1), "application/pdf")}).status_code == 401


def test_list_get_download_patch_delete(client, user, upload):
    h = user["headers"]
    f = upload(h)
    assert f["expires_at"] is None  # account files are kept
    assert any(x["id"] == f["id"] for x in client.get("/api/files?scope=recent", headers=h).json()["files"])
    assert client.get(f"/api/files/{f['id']}", headers=h).json()["name"] == "doc.pdf"
    data = client.get(f"/api/files/{f['id']}/download", headers=h).content
    assert data.startswith(b"%PDF")

    r = client.patch(f"/api/files/{f['id']}", headers=h, json={"name": "contract", "is_saved": True})
    assert r.json()["name"] == "contract.pdf" and r.json()["is_saved"]
    assert any(x["id"] == f["id"] for x in client.get("/api/files?scope=saved", headers=h).json()["files"])

    assert client.delete(f"/api/files/{f['id']}", headers=h).status_code == 204
    assert client.get(f"/api/files/{f['id']}", headers=h).status_code == 404


def test_guest_cannot_save(client, guest, upload):
    f = upload(guest)
    assert client.patch(f"/api/files/{f['id']}", headers=guest, json={"is_saved": True}).status_code == 401


def test_files_are_isolated(client, guest, upload, user):
    f = upload(guest)
    other = {"X-Guest-Token": uuid.uuid4().hex}
    assert client.get(f"/api/files/{f['id']}", headers=other).status_code == 404
    assert client.get(f"/api/files/{f['id']}/download", headers=user["headers"]).status_code == 404


def test_upload_version(client, guest, upload):
    f = upload(guest)
    r = client.post(f"/api/files/{f['id']}/versions", headers=guest, files={"file": ("x.pdf", make_pdf(1), "application/pdf")})
    assert r.status_code == 201 and r.json()["parent_id"] == f["id"]
    r = client.post(f"/api/files/{f['id']}/versions", headers=guest, files={"file": ("x.png", make_png(), "image/png")})
    assert r.status_code == 415


def test_inspect(client, guest, upload):
    f = upload(guest)
    base = f"/api/files/{f['id']}/inspect"
    info = client.get(f"{base}/info", headers=guest).json()["result"]
    assert info["page_count"] == 3
    spans = client.get(f"{base}/spans?page=0", headers=guest).json()["result"]
    assert any("Hello PDFDesk" in s["text"] for s in spans)
    images = client.get(f"{base}/images?page=0", headers=guest).json()["result"]
    assert len(images) == 1
    for what in ("links", "outline", "attachments", "annotations", "fields"):
        assert client.get(f"{base}/{what}?page=0", headers=guest).status_code == 200
    hits = client.get(f"{base}/search?q=PDFDesk", headers=guest).json()["result"]
    assert len(hits) == 3
    assert client.get(f"{base}/nope", headers=guest).status_code == 400
    assert client.get(f"{base}/spans?page=99", headers=guest).status_code == 400


def test_guest_files_expire(client, guest, upload):
    f = upload(guest)
    with SessionLocal() as db:
        row = db.get(StoredFile, f["id"])
        row.expires_at = datetime.now(timezone.utc) - timedelta(minutes=1)
        db.commit()
    assert cleanup_expired_files() >= 1
    assert client.get(f"/api/files/{f['id']}", headers=guest).status_code == 404


def test_encrypted_pdf_friendly_error(client, guest, upload, run):
    doc = fitz.open(stream=make_pdf(1))
    data = doc.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="u", owner_pw="o")
    f = upload(guest, data)
    job = run(guest, "rotate", [f["id"]], {"angle": 90}, expect="failed")
    assert "password" in job["error"].lower()
